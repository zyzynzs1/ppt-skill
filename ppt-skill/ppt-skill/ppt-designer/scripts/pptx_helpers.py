"""
ppt-designer 辅助函数库
基于 python-pptx，一键应用 design-tokens.md 中的设计令牌。
依赖: pip install python-pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------- 调色板 ----------
INK      = RGBColor(0x12, 0x26, 0x3F)  # 正文/深标题
NAVY     = RGBColor(0x0F, 0x2B, 0x52)  # 页标题/深色底
PRIMARY  = RGBColor(0x00, 0x68, 0xB7)  # 主蓝：强调/徽章/箭头
STEEL    = RGBColor(0x4A, 0x6B, 0x99)  # 次级标题/图表
MIST     = RGBColor(0x9F, 0xB8, 0xD8)
LINE     = RGBColor(0xD5, 0xE3, 0xF0)  # 描边/分隔线
TINT     = RGBColor(0xE3, 0xF0, 0xFB)
BG       = RGBColor(0xEF, 0xF5, 0xFB)  # 整页底色
CARD     = RGBColor(0xF4, 0xF9, 0xFF)  # 卡片底
GOLD     = RGBColor(0xC9, 0x98, 0x3B)  # 强调金
GOLD_DK  = RGBColor(0x8A, 0x5B, 0x14)
GOLD_LT  = RGBColor(0xEA, 0xD3, 0xAE)
GOLD_BG  = RGBColor(0xFD, 0xF9, 0xF0)
CYAN     = RGBColor(0x37, 0xB6, 0xE6)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"
PAGE_W, PAGE_H = Inches(10), Inches(5.625)

# ---------- 基础 ----------
def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = PAGE_W, PAGE_H
    return prs

def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])

def _fill(shape, color):
    shape.fill.solid(); shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False

def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, line=None):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(sp, color)
    if line:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    return sp

def text(slide, x, y, w, h, content, size=9, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """content: str 或 [(txt,size,color,bold), ...] 多段"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    items = content if isinstance(content, list) else [(content, size, color, bold)]
    for i, (t, s, c, b) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run(); run.text = t
        f = run.font; f.size = Pt(s); f.bold = b; f.color.rgb = c; f.name = FONT
        run._r.rPr.rFonts.set(qn('a:ea'), FONT)  # 中文字体
    return tb

def bullets(slide, x, y, w, h, items, size=9, color=INK, gap=4):
    """• 开头的要点列表，每条一行"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.2
        run = p.add_run(); run.text = "• " + it
        f = run.font; f.size = Pt(size); f.color.rgb = color; f.name = FONT
        run._r.rPr.rFonts.set(qn('a:ea'), FONT)
    return tb

# ---------- 页面骨架 ----------
def page_skeleton(slide, title, footer_left, page_no, lead=None):
    """内容页标准骨架：底色 + 标题组(+导语) + 页脚。返回内容区起点 y。"""
    rect(slide, 0, 0, 10, 5.625, BG)                       # 整页底色
    rect(slide, 0.375, 0.344, 0.0625, 0.266, PRIMARY)      # 标题装饰竖条
    text(slide, 0.523, 0.30, 6.5, 0.34, title, 17, NAVY, True)
    y = 0.85
    if lead:
        text(slide, 0.54, 0.85, 9.0, 0.36, lead, 9, INK)
        y = 1.25
    text(slide, 0.48, 5.44, 3.5, 0.15, footer_left, 9, STEEL)
    text(slide, 9.16, 5.44, 0.4, 0.15, f"{page_no:02d}", 9, STEEL, align=PP_ALIGN.RIGHT)
    return y

# ---------- 组件 ----------
def badge(slide, x, y, label, d=0.38, style="A", size=11):
    """编号徽章。style A: 主蓝底白字；style B: 白底蓝框蓝字"""
    sp = rect(slide, x, y, d, d, PRIMARY if style == "A" else WHITE,
              MSO_SHAPE.OVAL, line=None if style == "A" else PRIMARY)
    text(slide, x, y + d/2 - 0.09, d, 0.2, label, size,
         WHITE if style == "A" else PRIMARY, True, PP_ALIGN.CENTER)
    return sp

def card(slide, x, y, w, h, title, en=None, points=None, fill=CARD):
    """信息卡片：白/浅蓝底 + 标题 + 英文副题 + 要点列表"""
    rect(slide, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, line=LINE)
    pad = 0.12
    text(slide, x + pad, y + pad, w - 2*pad, 0.24, title, 12, NAVY, True)
    yy = y + pad + 0.26
    if en:
        text(slide, x + pad, yy, w - 2*pad, 0.15, en, 7.5, STEEL); yy += 0.22
    if points:
        bullets(slide, x + pad, yy + 0.06, w - 2*pad, h - (yy - y) - pad, points, 8)
    return

def card_row(slide, y, titles, ens=None, points_list=None, h=2.0, n=None):
    """等宽卡片横排（3~4 列），自动算列宽，同排等高"""
    n = n or len(titles)
    gap = 0.18
    w = (9.0 - (n - 1) * gap) / n
    for i in range(n):
        card(slide, 0.5 + i * (w + gap), y, w, h, titles[i],
             ens[i] if ens else None,
             points_list[i] if points_list else None)

def arrow(slide, x1, y1, x2, y2, color=PRIMARY, wpt=1.5):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = color; conn.line.width = Pt(wpt)
    conn.line._get_or_add_ln().append(
        conn.line._get_or_add_ln()._new_tailEnd())
    conn.line._get_or_add_ln().tailEnd.set('type', 'arrow')
    return conn

def flow_chain(slide, y, steps, x0=0.5, x1=9.5, h=1.0, badge_style="A"):
    """横向步骤链：[(标题, 说明), ...]，节点间自动连箭头"""
    n = len(steps)
    gap = 0.25
    w = (x1 - x0 - (n - 1) * gap) / n
    for i, (t, desc) in enumerate(steps):
        x = x0 + i * (w + gap)
        badge(slide, x, y, f"{i+1:02d}", style=badge_style)
        text(slide, x + 0.5, y + 0.02, w - 0.5, 0.2, t, 10, NAVY, True)
        text(slide, x, y + 0.45, w, h - 0.45, desc, 8, STEEL)
        if i < n - 1:
            arrow(slide, x + w + 0.03, y + 0.19, x + w + gap - 0.03, y + 0.19)

def kpi_row(slide, y, items, x0=0.5, x1=9.5, num_color=PRIMARY):
    """大数字指标组：[(数字, 单位, 说明), ...]"""
    n = len(items)
    w = (x1 - x0) / n
    for i, (num, unit, desc) in enumerate(items):
        x = x0 + i * w
        if i > 0:
            rect(slide, x, y + 0.1, 0.008, 0.8, LINE)
        text(slide, x + 0.1, y, w - 0.2, 0.5,
             [(num, 32, num_color, True), (" " + unit, 12, STEEL, False)])
        text(slide, x + 0.1, y + 0.55, w - 0.2, 0.3, desc, 8, STEEL)

def conclusion_bar(slide, y, sentence, h=0.35):
    """通栏浅金结论条"""
    rect(slide, 0.5, y, 9.0, h, GOLD_BG, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(slide, 0.62, y + 0.08, 0.04, h - 0.16, GOLD)
    text(slide, 0.78, y + 0.06, 8.6, h - 0.1, sentence, 10, INK, True,
         anchor=MSO_ANCHOR.MIDDLE)

def section_band(slide, y, label, x=0.5, w=9.0, h=0.3):
    """深蓝阶段带（流程分阶段用）"""
    rect(slide, x, y, w, h, NAVY)
    text(slide, x + 0.15, y + 0.04, w - 0.3, h - 0.06, label, 10, WHITE, True,
         anchor=MSO_ANCHOR.MIDDLE)

# ---------- 示例：生成一页三卡片内容页 ----------
if __name__ == "__main__":
    prs = new_deck()
    s = blank(prs)
    page_skeleton(s, "本页结论式标题写在这里", "主题 · 汇报方案", 3,
                  lead="标题下可放一行导语，9pt，概括本页核心逻辑。")
    card_row(s, 1.35,
             ["能力一", "能力二", "能力三"],
             ["CAPABILITY ONE", "CAPABILITY TWO", "CAPABILITY THREE"],
             [["要点 A，一行一条", "要点 B", "要点 C"],
              ["要点 A", "要点 B", "要点 C"],
              ["要点 A", "要点 B", "要点 C"]],
             h=2.2)
    conclusion_bar(s, 4.6, "一句话收束本页观点，金色点睛，全稿每页至多一处。")
    prs.save("demo.pptx")
    print("demo.pptx 已生成")
