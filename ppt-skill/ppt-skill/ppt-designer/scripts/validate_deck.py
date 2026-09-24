#!/usr/bin/env python3
"""ppt-designer 自动质检器
用法: python validate_deck.py 文件.pptx
检查: 调色板外颜色 / 字体混用 / 页脚缺失 / 页面尺寸 / 文本框出界
"""
import sys
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE

ALLOWED_COLORS = {
    "12263F", "0F2B52", "0068B7", "4A6B99", "9FB8D8", "D5E3F0", "E3F0FB",
    "EFF5FB", "F4F9FF", "C9983B", "8A5B14", "EAD3AE", "FDF9F0", "37B6E6",
    "FFFFFF",
    # 备选色板 B · 墨绿商务
    "1A2E1F", "152A1E", "2F6B46", "527A5E", "A8C4B0", "D8E6DC",
    "E6F2EA", "F1F7F2", "F7FBF8",
    # 备选色板 C · 石墨炽橙
    "1F2328", "14181D", "E8722A", "5A6B7E", "B8C4D0", "DCE3EA",
    "EDF1F5", "F4F6F9", "FAFBFD",
}
ALLOWED_FONTS = {"微软雅黑", "Microsoft YaHei", "Arial", None}
EMU_IN = 914400


def run_color(obj, where, errors):
    try:
        if obj.fill.type == MSO_FILL_TYPE.SOLID:
            c = obj.fill.fore_color
            if c.type is not None and str(c.rgb) not in ALLOWED_COLORS:
                errors.append(f"[颜色越界] {where}: #{c.rgb}")
    except Exception:
        pass


def check(path):
    prs = Presentation(path)
    errors, warnings = [], []
    # 页面尺寸
    if abs(prs.slide_width - 10 * EMU_IN) > 1000 or abs(prs.slide_height - 5.625 * EMU_IN) > 1000:
        errors.append(
            f"[尺寸] 非 10x5.625in 16:9 画布: "
            f"{prs.slide_width / EMU_IN:.2f}x{prs.slide_height / EMU_IN:.2f}"
        )
    for i, slide in enumerate(prs.slides, 1):
        has_footer = False
        for sh in slide.shapes:
            where = f"第{i}页 {sh.shape_type}({sh.shape_id})"
            run_color(sh, where, errors)
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        # 字体混用
                        if r.font.name not in ALLOWED_FONTS:
                            errors.append(f"[字体混用] {where}: {r.font.name}")
                        # 文字颜色越界
                        try:
                            if r.font.color and r.font.color.type is not None \
                                    and str(r.font.color.rgb) not in ALLOWED_COLORS:
                                errors.append(f"[文字颜色越界] {where}: #{r.font.color.rgb}")
                        except Exception:
                            pass
            # 出界 / 页脚检测
            try:
                if sh.left is not None and sh.top is not None:
                    if (sh.left < -10000 or sh.top < -10000
                            or sh.left + sh.width > prs.slide_width + 10000
                            or sh.top + sh.height > prs.slide_height + 10000):
                        warnings.append(f"[出界] {where}: 超出画布")
                    if sh.top / EMU_IN > 5.35 and sh.has_text_frame \
                            and sh.text_frame.text.strip():
                        has_footer = True
            except Exception:
                pass
        if i > 1 and not has_footer:
            warnings.append(f"[页脚缺失] 第{i}页底部未检测到文本")
    print(f"== 校验结果: {len(errors)} 错误 / {len(warnings)} 警告 ==")
    for e in errors:
        print("  ✗", e)
    for w in warnings:
        print("  ⚠", w)
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python validate_deck.py 文件.pptx")
        sys.exit(2)
    check(sys.argv[1])
